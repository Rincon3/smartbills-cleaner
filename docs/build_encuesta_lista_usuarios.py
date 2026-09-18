from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path


root = Path("/Users/miguelrinconclavijo/Desktop/study/Maestria en desarrollo IA/trabajo de grado maestria/smartbills")
out_dir = root / "outputs" / "encuesta-usuarios-2026-05-23"
out_dir.mkdir(parents=True, exist_ok=True)
out_file = out_dir / "encuesta_smartbills_lista_para_llenar.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "Encuesta"
ws2 = wb.create_sheet("Respuestas")
ws3 = wb.create_sheet("Codigos")

header_fill = PatternFill("solid", fgColor="0F8B7D")
section_fill = PatternFill("solid", fgColor="DFF5F1")
sub_fill = PatternFill("solid", fgColor="EAF3FF")
thin = Side(style="thin", color="D5E2DF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_range(sheet, cell_range, fill=None, bold=False, size=11, center=False):
    for row in sheet[cell_range]:
        for cell in row:
            cell.border = border
            cell.font = Font(bold=bold, size=size, color="FFFFFF" if fill == header_fill else "16322F")
            if fill:
                cell.fill = fill
            if center:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(vertical="center", wrap_text=True)


encuesta_rows = [
    ["ENCUESTA A USUARIOS - SMARTBILLS CLEANER", "", "", "", "", ""],
    ["Objetivo", "Conocer hábitos, necesidades y percepción de usabilidad en la gestión de facturas digitales.", "", "", "", ""],
    ["", "", "", "", "", ""],
    ["DATOS DEL PARTICIPANTE", "", "", "", "", ""],
    ["ID participante", "", "Fecha", "", "Entrevistador", ""],
    ["Edad", "", "Ocupación", "", "Canal", ""],
    ["Nivel de experiencia digital", "", "Tipo de usuario", "", "Consentimiento", ""],
    ["", "", "", "", "", ""],
    ["SECCIÓN 1. CONTEXTO DE USO", "", "", "", "", ""],
    ["P1", "¿Con qué frecuencia gestionas facturas o documentos similares?", "Nunca", "Ocasionalmente", "Frecuentemente", "Diariamente"],
    ["Respuesta P1", "", "", "", "", ""],
    ["P2", "¿Dónde sueles almacenar tus facturas?", "Correo", "Carpetas PC", "Nube", "WhatsApp", "Otro"],
    ["Respuesta P2", "", "", "", "", ""],
    ["P3", "¿Qué tan difícil te resulta encontrar una factura cuando la necesitas?", "1 Muy fácil", "2 Fácil", "3 Media", "4 Difícil", "5 Muy difícil"],
    ["Respuesta P3", "", "", "", "", ""],
    ["", "", "", "", "", ""],
    ["SECCIÓN 2. PROBLEMAS ACTUALES", "", "", "", "", ""],
    ["P4", "¿Has tenido errores al revisar datos como proveedor, fecha o valor?", "Sí", "No", "", "", ""],
    ["Respuesta P4", "", "", "", "", ""],
    ["P5", "¿Qué problema ocurre con mayor frecuencia?", "Pérdida de tiempo", "Documentos desordenados", "Errores manuales", "Dificultad para buscar", "Otro"],
    ["Respuesta P5", "", "", "", "", ""],
    ["P6", "Describe brevemente el principal problema que enfrentas hoy", "", "", "", "", ""],
    ["Respuesta P6", "", "", "", "", ""],
    ["", "", "", "", "", ""],
    ["SECCIÓN 3. EXPECTATIVAS SOBRE LA APLICACIÓN", "", "", "", "", ""],
    ["P7", "¿Qué tan útil te parece una aplicación que extraiga automáticamente datos de facturas?", "1 Nada útil", "2 Poco útil", "3 Media", "4 Útil", "5 Muy útil"],
    ["Respuesta P7", "", "", "", "", ""],
    ["P8", "¿Qué tan importante sería poder corregir datos antes de guardarlos?", "1 Nada importante", "2 Poco importante", "3 Media", "4 Importante", "5 Muy importante"],
    ["Respuesta P8", "", "", "", "", ""],
    ["P9", "¿Qué tan útil sería consultar facturas con preguntas en lenguaje natural?", "1 Nada útil", "2 Poco útil", "3 Media", "4 Útil", "5 Muy útil"],
    ["Respuesta P9", "", "", "", "", ""],
    ["P10", "¿Qué módulo te parecería más valioso?", "Carga de facturas", "Listado y filtros", "Validación de datos", "Consultas inteligentes", "Dashboard"],
    ["Respuesta P10", "", "", "", "", ""],
    ["", "", "", "", "", ""],
    ["SECCIÓN 4. USABILIDAD ESPERADA", "", "", "", "", ""],
    ["P11", "¿Qué tan fácil debería ser usar la aplicación?", "1 Muy difícil", "2 Difícil", "3 Media", "4 Fácil", "5 Muy fácil"],
    ["Respuesta P11", "", "", "", "", ""],
    ["P12", "¿Qué tan importante es que el sistema muestre el estado del procesamiento?", "1 Nada importante", "2 Poco importante", "3 Media", "4 Importante", "5 Muy importante"],
    ["Respuesta P12", "", "", "", "", ""],
    ["P13", "¿Prefieres una interfaz simple o una más completa con más información?", "1 Muy simple", "2 Simple", "3 Intermedia", "4 Completa", "5 Muy completa"],
    ["Respuesta P13", "", "", "", "", ""],
    ["", "", "", "", "", ""],
    ["SECCIÓN 5. PERCEPCIÓN DE LA INTERFAZ", "", "", "", "", ""],
    ["P14", "¿La distribución general de la interfaz te parece ordenada y fácil de seguir?", "1 Muy en desacuerdo", "2 En desacuerdo", "3 Neutral", "4 De acuerdo", "5 Muy de acuerdo"],
    ["Respuesta P14", "", "", "", "", ""],
    ["P15", "¿La interfaz te ayuda a entender rápidamente de qué trata la aplicación?", "1 Muy en desacuerdo", "2 En desacuerdo", "3 Neutral", "4 De acuerdo", "5 Muy de acuerdo"],
    ["Respuesta P15", "", "", "", "", ""],
    ["P16", "¿Los módulos o secciones visibles se entienden con claridad?", "1 Muy en desacuerdo", "2 En desacuerdo", "3 Neutral", "4 De acuerdo", "5 Muy de acuerdo"],
    ["Respuesta P16", "", "", "", "", ""],
    ["P17", "¿La organización visual transmite confianza para usar la aplicación?", "1 Muy en desacuerdo", "2 En desacuerdo", "3 Neutral", "4 De acuerdo", "5 Muy de acuerdo"],
    ["Respuesta P17", "", "", "", "", ""],
    ["P18", "¿La distribución de botones, paneles y opciones parece lógica?", "1 Muy en desacuerdo", "2 En desacuerdo", "3 Neutral", "4 De acuerdo", "5 Muy de acuerdo"],
    ["Respuesta P18", "", "", "", "", ""],
    ["P19", "¿La interfaz parece adecuada para revisar y gestionar facturas sin confusión?", "1 Muy en desacuerdo", "2 En desacuerdo", "3 Neutral", "4 De acuerdo", "5 Muy de acuerdo"],
    ["Respuesta P19", "", "", "", "", ""],
    ["P20", "Observación final del participante sobre la interfaz", "", "", "", "", ""],
    ["Respuesta P20", "", "", "", "", ""],
]

for row in encuesta_rows:
    ws.append(row)

ws.merge_cells("A1:F1")
ws.merge_cells("B2:F2")
ws.merge_cells("A4:F4")
ws.merge_cells("A9:F9")
ws.merge_cells("A17:F17")
ws.merge_cells("A25:F25")
ws.merge_cells("A35:F35")
ws.merge_cells("A44:F44")

style_range(ws, "A1:F1", fill=header_fill, bold=True, size=14, center=True)
style_range(ws, "A4:F4", fill=section_fill, bold=True, size=12)
style_range(ws, "A9:F9", fill=section_fill, bold=True, size=12)
style_range(ws, "A17:F17", fill=section_fill, bold=True, size=12)
style_range(ws, "A25:F25", fill=section_fill, bold=True, size=12)
style_range(ws, "A35:F35", fill=section_fill, bold=True, size=12)
style_range(ws, "A44:F44", fill=section_fill, bold=True, size=12)
style_range(ws, "A2:F56")

for r in [10,12,14,18,20,22,26,28,30,32,36,38,40,45,47,49,51,53,55]:
    style_range(ws, f"A{r}:F{r}", fill=sub_fill, bold=True)

for r in [11,13,15,19,21,23,27,29,31,33,37,39,41,43,46,48,50,52,54,56]:
    for cell in ws[f"A{r}:F{r}"][0]:
        cell.fill = PatternFill("solid", fgColor="FFFDF4")

widths = {1: 18, 2: 42, 3: 18, 4: 18, 5: 18, 6: 18}
for col, width in widths.items():
    ws.column_dimensions[get_column_letter(col)].width = width

resp_rows = [
    ["id_participante", "fecha", "pregunta", "codigo", "tipo_respuesta", "respuesta_cerrada", "valor_numerico", "respuesta_abierta", "perfil", "tipo_usuario", "observaciones"],
]
for code, pregunta, tipo in [
    ("P1", "Frecuencia de gestión de facturas", "cerrada"),
    ("P2", "Lugar de almacenamiento", "cerrada"),
    ("P3", "Dificultad para encontrar una factura", "escala_1_5"),
    ("P4", "Errores al revisar datos", "cerrada"),
    ("P5", "Problema más frecuente", "cerrada"),
    ("P6", "Principal problema actual", "abierta"),
    ("P7", "Utilidad de la automatización", "escala_1_5"),
    ("P8", "Importancia de corregir datos", "escala_1_5"),
    ("P9", "Utilidad de consultas en lenguaje natural", "escala_1_5"),
    ("P10", "Módulo más valioso", "cerrada"),
    ("P11", "Facilidad esperada del sistema", "escala_1_5"),
    ("P12", "Importancia de retroalimentación del sistema", "escala_1_5"),
    ("P13", "Preferencia de complejidad de interfaz", "escala_1_5"),
    ("P14", "Orden y distribución general de la interfaz", "escala_1_5"),
    ("P15", "Comprensión rápida del propósito de la aplicación", "escala_1_5"),
    ("P16", "Claridad de los módulos visibles", "escala_1_5"),
    ("P17", "Confianza transmitida por la organización visual", "escala_1_5"),
    ("P18", "Lógica en la distribución de botones y paneles", "escala_1_5"),
    ("P19", "Adecuación visual para gestionar facturas sin confusión", "escala_1_5"),
    ("P20", "Observación final sobre la interfaz", "abierta"),
]:
    resp_rows.append(["P01", "2026-05-23", pregunta, code, tipo, "", "", "", "", "", ""])

for row in resp_rows:
    ws2.append(row)

style_range(ws2, f"A1:K{len(resp_rows)}")
style_range(ws2, "A1:K1", fill=header_fill, bold=True, center=True)
for i, width in enumerate([16, 14, 34, 10, 16, 18, 14, 28, 14, 14, 18], start=1):
    ws2.column_dimensions[get_column_letter(i)].width = width

codigo_rows = [
    ["codigo", "seccion", "tipo", "pregunta", "opciones", "uso_estadistico"],
    ["P1", "Contexto", "cerrada", "¿Con qué frecuencia gestionas facturas o documentos similares?", "Nunca; Ocasionalmente; Frecuentemente; Diariamente", "Segmentación por frecuencia"],
    ["P2", "Contexto", "cerrada", "¿Dónde sueles almacenar tus facturas?", "Correo; Carpetas PC; Nube; WhatsApp; Otro", "Distribución por medio de almacenamiento"],
    ["P3", "Contexto", "escala_1_5", "¿Qué tan difícil te resulta encontrar una factura cuando la necesitas?", "1 a 5", "Promedio de fricción"],
    ["P4", "Problemas", "cerrada", "¿Has tenido errores al revisar datos como proveedor, fecha o valor?", "Sí; No", "Incidencia de errores"],
    ["P5", "Problemas", "cerrada", "¿Qué problema ocurre con mayor frecuencia?", "Pérdida de tiempo; Desorden; Errores; Búsqueda; Otro", "Problema dominante"],
    ["P6", "Problemas", "abierta", "Describe brevemente el principal problema que enfrentas hoy", "Texto libre", "Hallazgos cualitativos"],
    ["P7", "Expectativas", "escala_1_5", "Utilidad de automatizar extracción", "1 a 5", "Promedio de utilidad"],
    ["P8", "Expectativas", "escala_1_5", "Importancia de corregir antes de guardar", "1 a 5", "Promedio de control/confianza"],
    ["P9", "Expectativas", "escala_1_5", "Utilidad de consultas en lenguaje natural", "1 a 5", "Valor percibido de IA"],
    ["P10", "Expectativas", "cerrada", "Módulo más valioso", "Carga; Listado; Validación; Consultas; Dashboard", "Priorización funcional"],
    ["P11", "Usabilidad", "escala_1_5", "Facilidad esperada", "1 a 5", "Meta de usabilidad"],
    ["P12", "Usabilidad", "escala_1_5", "Importancia del estado del proceso", "1 a 5", "Necesidad de retroalimentación"],
    ["P13", "Usabilidad", "escala_1_5", "Preferencia por simplicidad/completitud", "1 a 5", "Preferencia de complejidad"],
    ["P14", "Interfaz", "escala_1_5", "¿La distribución general de la interfaz te parece ordenada y fácil de seguir?", "1 a 5", "Orden visual percibido"],
    ["P15", "Interfaz", "escala_1_5", "¿La interfaz te ayuda a entender rápidamente de qué trata la aplicación?", "1 a 5", "Comprensión inicial de propósito"],
    ["P16", "Interfaz", "escala_1_5", "¿Los módulos o secciones visibles se entienden con claridad?", "1 a 5", "Claridad estructural"],
    ["P17", "Interfaz", "escala_1_5", "¿La organización visual transmite confianza para usar la aplicación?", "1 a 5", "Confianza visual"],
    ["P18", "Interfaz", "escala_1_5", "¿La distribución de botones, paneles y opciones parece lógica?", "1 a 5", "Lógica de layout"],
    ["P19", "Interfaz", "escala_1_5", "¿La interfaz parece adecuada para revisar y gestionar facturas sin confusión?", "1 a 5", "Adecuación funcional visual"],
    ["P20", "Interfaz", "abierta", "Observación final del participante sobre la interfaz", "Texto libre", "Comentario cualitativo de interfaz"],
]
for row in codigo_rows:
    ws3.append(row)
style_range(ws3, f"A1:F{len(codigo_rows)}")
style_range(ws3, "A1:F1", fill=header_fill, bold=True, center=True)
for i, width in enumerate([10, 14, 14, 42, 34, 24], start=1):
    ws3.column_dimensions[get_column_letter(i)].width = width

wb.save(out_file)
print(out_file)
